"""Formato de columnas de los reportes (saca-tu-garra, fondeo-estable) según governance/tasks/tarea.md."""

from datetime import date, datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from reportes.comun import ejecutor
from reportes.comun.fechas import Cortes
from reportes.mensuales.piero.r02_fondeo_estable import REPORTE as FONDEO
from reportes.mensuales.piero.r06_saca_tu_garra import REPORTE as SACA


def test_saca_tu_garra_columnas_encabezados_y_reemplazos(tmp_path):
    df = pd.DataFrame({
        "HASEOPER": ["U001", "U002", "U003"], "VAR_VIGENTE": [1500.5, 0, -20.0], "PRODUCTIVDAD": [3, 0, 0],
        "RatioRecuperacion0_30": [0.8, 0.0, None], "RatioRecuperacion1_30": [0, 0.25, 0.5], "OTRA": [1, 2, 3],
    })
    ruta = ejecutor.exportar_excel(SACA, [df], Cortes(date(2026, 9, 30)), tmp_path)
    ws = load_workbook(ruta).active
    assert [c.value for c in ws[1]] == ["Usuario", "Var. Saldo Vigente", "Productividad", "Efectividad -30 a 0", "Efectividad 1 a 30"]
    filas = [[c.value for c in f] for f in ws.iter_rows(min_row=2)]
    assert filas[0] == ["U001", 1500.5, 3, 0.8, "NULL"]      # ratio 0 => NULL
    assert filas[1] == ["U002", "-", "-", "NULL", 0.25]        # var y productividad 0 => "-"; ratio 0 => NULL
    assert filas[2] == ["U003", -20.0, "-", "NULL", 0.5]       # ratio nulo => NULL
    assert ruta.name == "Base Saca tu Garra_20260930.xlsx"


def test_saca_tu_garra_acepta_productividad_bien_escrita(tmp_path):
    df = pd.DataFrame({"HASEOPER": ["U1"], "VAR_VIGENTE": [1], "PRODUCTIVIDAD": [0],
                       "RatioRecuperacion0_30": [1], "RatioRecuperacion1_30": [1]})
    ws = load_workbook(ejecutor.exportar_excel(SACA, [df], Cortes(date(2026, 9, 30)), tmp_path)).active
    assert ws["C2"].value == "-"


def test_fondeo_estable_fecha_dd_mm_aaaa_y_encabezados(tmp_path):
    df = pd.DataFrame({"FECHA": [datetime(2026, 9, 30)] * 2, "RDESMAT": ["MATRIZ A", "MATRIZ B"], "HSALFESI": [1234.5, 99.0]})
    ruta = ejecutor.exportar_excel(FONDEO, [df], Cortes(date(2026, 9, 30)), tmp_path)
    ws = load_workbook(ruta).active
    assert [c.value for c in ws[1]] == ["FECHA", "Matriz", "Saldo Fondeo Estable"]
    assert ws["A2"].value.date() == date(2026, 9, 30) if isinstance(ws["A2"].value, datetime) else ws["A2"].value == date(2026, 9, 30)
    assert ws["A2"].number_format == "dd/mm/yyyy"
    assert ruta.name == "Saldo_FondeoEstable_20260930.xlsx"


def test_columna_faltante_es_error_de_datos_claro(tmp_path):
    with pytest.raises(ejecutor.ErrorDatos, match="no trae la columna"):
        ejecutor.exportar_excel(FONDEO, [pd.DataFrame({"FECHA": [date(2026, 9, 30)]})], Cortes(date(2026, 9, 30)), tmp_path)

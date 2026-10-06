"""Formato Excel de las salidas, copiado del legado (Cartera-Sin asignar-base.xlsm).

- `escribir_tabla`: datos como Tabla de Excel con estilo «TableStyleMedium2» y anchos de columna (hoja DATA_MIS_v2).
- `escribir_resumen_jerarquico`: resumen tipo tabla dinámica (hoja RESUMEN_v2): niveles con sangría, suma por nivel,
  formato `"S/" #,##0.00` y «Total general»; colores de la imagen Reporte_Temporal.jpg que se enviaba por correo.
- `guardar_libro`: escribe las hojas en un .xlsx; con `compartido=True` reutiliza el libro si ya existe (varios
  comandos alimentan el mismo archivo, como el «Datos Cierre» de Michael).

No agrega hojas de control ni metadatos: el libro contiene solo los datos del reporte.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ESTILO_TABLA = "TableStyleMedium2"
FORMATO_MONEDA = '"S/" #,##0.00'
EN_BLANCO = "(en blanco)"

# Colores del resumen (imagen Reporte_Temporal.jpg del legado)
AZUL_ENCABEZADO = "002060"
AZUL_TOTAL = "244062"
AZUL_NIVEL_1 = "DCE6F1"
AZUL_NIVEL_2 = "F2F6FA"
_FINO = Side(style="thin", color="000000")
_BORDE = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)


@dataclass(frozen=True)
class Resumen:
    """Hoja de resumen jerárquico calculada a partir del resultado `origen` (índice en la lista de resultados)."""

    nombre: str                   # nombre de la hoja (RESUMEN_v2)
    niveles: tuple[str, ...]      # columnas del resultado, de mayor a menor nivel
    valor: str                    # columna que se suma
    titulo_filas: str             # encabezado de la columna de etiquetas
    titulo_valor: str             # encabezado de la columna de valores
    origen: int = 0
    formato: str = FORMATO_MONEDA


# ---------------------------------------------------------------- tabla de datos
def _normalizar(df: pd.DataFrame) -> pd.DataFrame:
    """Columnas con nombre texto único; fechas sin hora si todas son medianoche."""
    df = df.copy()
    vistos: dict[str, int] = {}
    nombres = []
    for c in df.columns:
        n = "Columna" if c is None or (not isinstance(c, str) and pd.isna(c)) else str(c)
        vistos[n] = vistos.get(n, 0) + 1
        nombres.append(n if vistos[n] == 1 else f"{n}_{vistos[n]}")
    df.columns = nombres
    for c in nombres:
        if pd.api.types.is_datetime64_any_dtype(df[c]):
            serie = df[c].dropna()
            df[c] = df[c].dt.date if not serie.empty and (serie == serie.dt.normalize()).all() else df[c]
    return df


def _valor_excel(v):
    if v is None or (not isinstance(v, (str, bytes)) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item"):  # numpy escalar
        return v.item()
    return v


def _ajustar_anchos(ws, minimo: int = 8, maximo: int = 60) -> None:
    for col in ws.columns:
        letra = col[0].column_letter
        ancho = max((len(str(c.value)) for c in col[:300] if c.value is not None), default=minimo)
        ws.column_dimensions[letra].width = min(max(minimo, ancho + 2), maximo)


def nombre_tabla(nombre_hoja: str) -> str:
    """Nombre de Tabla de Excel válido: letras, números y _, sin empezar por número."""
    limpio = re.sub(r"\W", "_", nombre_hoja)
    return limpio if re.match(r"[A-Za-z_]", limpio) else f"T_{limpio}"


def escribir_tabla(ws, df: pd.DataFrame) -> None:
    """Encabezados + filas como Tabla de Excel (TableStyleMedium2), como DATA_MIS_v2 del legado."""
    df = _normalizar(df)
    ws.append(list(df.columns))
    for fila in df.itertuples(index=False, name=None):
        ws.append([_valor_excel(v) for v in fila])
    for idx, c in enumerate(df.columns, 1):
        if df[c].map(lambda v: isinstance(v, date) and not isinstance(v, datetime)).any():
            for celda in ws.iter_cols(min_col=idx, max_col=idx, min_row=2):
                for x in celda:
                    x.number_format = "yyyy-mm-dd"
    if len(df) >= 1 and len(df.columns) >= 1:  # Excel exige al menos una fila de datos para definir una Tabla
        ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
        tabla = Table(displayName=nombre_tabla(ws.title), ref=ref)
        tabla.tableStyleInfo = TableStyleInfo(name=ESTILO_TABLA, showRowStripes=True)
        ws.add_table(tabla)
    _ajustar_anchos(ws)


# ---------------------------------------------------------------- resumen jerárquico
def _etiqueta(v) -> str:
    return EN_BLANCO if v is None or (not isinstance(v, str) and pd.isna(v)) or str(v).strip() == "" else str(v)


def _clave_orden(texto: str) -> tuple[int, str]:
    return (1, "") if texto == EN_BLANCO else (0, texto.casefold())


def filas_jerarquicas(df: pd.DataFrame, niveles: tuple[str, ...], valor: str) -> list[tuple[int, str, float]]:
    """[(nivel, etiqueta, suma)] en orden de tabla dinámica: cada nivel ordenado A→Z, «(en blanco)» al final."""
    trabajo = df[list(niveles) + [valor]].copy()
    for n in niveles:
        trabajo[n] = trabajo[n].map(_etiqueta)
    trabajo[valor] = pd.to_numeric(trabajo[valor], errors="coerce").fillna(0.0)

    salida: list[tuple[int, str, float]] = []

    def recorrer(sub: pd.DataFrame, nivel: int) -> None:
        if nivel == len(niveles):
            return
        sumas = sub.groupby(niveles[nivel], sort=False)[valor].sum()
        for etiqueta in sorted(sumas.index, key=_clave_orden):
            salida.append((nivel, etiqueta, float(sumas[etiqueta])))
            recorrer(sub[sub[niveles[nivel]] == etiqueta], nivel + 1)

    recorrer(trabajo, 0)
    return salida


def escribir_resumen_jerarquico(ws, df: pd.DataFrame, r: Resumen) -> None:
    """Resumen tipo tabla dinámica en B2:C… (como RESUMEN_v2): sangría por nivel, suma, moneda y Total general."""
    filas = filas_jerarquicas(df, r.niveles, r.valor)
    total = float(pd.to_numeric(df[r.valor], errors="coerce").fillna(0.0).sum())
    ws.cell(2, 2, r.titulo_filas)
    ws.cell(2, 3, r.titulo_valor)
    for col in (2, 3):
        c = ws.cell(2, col)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL_ENCABEZADO)
        c.border = _BORDE
        c.alignment = Alignment(horizontal="center", vertical="center")
    ultimo = len(r.niveles) - 1
    for i, (nivel, etiqueta, suma) in enumerate(filas, start=3):
        b, c = ws.cell(i, 2, etiqueta), ws.cell(i, 3, suma)
        b.alignment = Alignment(indent=nivel)
        c.number_format = r.formato
        relleno = AZUL_NIVEL_1 if nivel <= 1 else AZUL_NIVEL_2 if nivel < ultimo else None
        for x in (b, c):
            x.border = _BORDE
            x.font = Font(name="Calibri", size=11, bold=nivel < ultimo)
            if relleno:
                x.fill = PatternFill("solid", fgColor=relleno)
    fila_total = len(filas) + 3
    ws.cell(fila_total, 2, "Total general")
    ws.cell(fila_total, 3, total).number_format = r.formato
    for col in (2, 3):
        c = ws.cell(fila_total, col)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL_TOTAL)
        c.border = _BORDE
    ws.column_dimensions["B"].width = 34.9
    ws.column_dimensions["C"].width = 19.9
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------- libro
def _nombre_hoja_unico(nombre: str, usados: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", "_", nombre)[:31] or "Hoja"
    cand, n = base, 2
    while cand in usados:
        cand = f"{base[:28]}_{n}"
        n += 1
    usados.add(cand)
    return cand


def guardar_libro(ruta: Path, hojas: list[tuple[str, pd.DataFrame]], resumenes: list[tuple[Resumen, pd.DataFrame]] = (),
                  compartido: bool = False) -> Path:
    """Escribe las hojas (tabla) y los resúmenes en `ruta`.

    `compartido=True`: si el libro ya existe se conserva y solo se reemplazan las hojas con el mismo nombre.
    """
    ruta.parent.mkdir(parents=True, exist_ok=True)
    if compartido and ruta.exists():
        wb = load_workbook(ruta)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    usados = set(wb.sheetnames) if compartido else set()
    for nombre, df in hojas:
        nombre = nombre if compartido else _nombre_hoja_unico(nombre, usados)
        nombre = re.sub(r"[\[\]\*\?/\\:]", "_", nombre)[:31]
        if nombre in wb.sheetnames:
            del wb[nombre]
        escribir_tabla(wb.create_sheet(nombre), df)
    for resumen, df in resumenes:
        nombre = resumen.nombre if compartido else _nombre_hoja_unico(resumen.nombre, usados)
        if nombre in wb.sheetnames:
            del wb[nombre]
        escribir_resumen_jerarquico(wb.create_sheet(nombre), df, resumen)
    if not wb.sheetnames:
        wb.create_sheet("Datos")
    wb.save(ruta)
    return ruta

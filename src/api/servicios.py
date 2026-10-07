"""Casos de uso de la interfaz web sobre el motor de `reportes` (catálogo, tablas, verificación, solicitud, archivos, envío).

No reimplementa reglas: llama a las mismas funciones que usa la consola, así la web y `main.py` no pueden divergir.
Las validaciones lanzan `PeticionInvalida` (422) o `NoEncontrado` (404); la capa HTTP solo las traduce.
"""

from __future__ import annotations

import json
import time
from datetime import date, datetime
from decimal import Decimal
from functools import lru_cache
from importlib import import_module
from pathlib import Path

from reportes.comun.ejecutor import ReporteLote
from reportes.comun.fechas import Cortes, resolver_corte
from reportes.config import BASES_DE, DIR_INPUTS, DIR_OUTPUTS, DRIVER_ODBC, SERVIDORES, ConfiguracionError
from reportes.db import leer_sql
from reportes.registro import DIA_ANTERIOR_SIMPLE, REPORTES, Reporte, carpeta_salida, ordenados
from reportes.reglas_fecha import regla_de
from reportes.tablas import tablas_de
from reportes.verificacion import Estado, mensaje_solicitud, nombre_resuelto, verificar_reporte

from . import esquemas as e

FILAS_VISTA_PREVIA = 200
MAX_TEXTO = 200_000
# Reportes de Erick que reciben el mes (AAAA-MM) en lugar de la fecha de corte.
USAN_MES = frozenset({"bancarizados-producto", "indicadores-clientes"})
TIPOS = {".xlsx": "excel", ".xlsm": "excel", ".jpg": "imagen", ".jpeg": "imagen", ".png": "imagen", ".txt": "texto", ".csv": "texto", ".json": "texto"}


class PeticionInvalida(ValueError):
    """La petición no cumple una regla del reporte (fecha, confirmación, correo…)."""


class NoEncontrado(LookupError):
    pass


# ---------------------------------------------------------------- catálogo

def reporte(nombre: str) -> Reporte:
    if nombre not in REPORTES:
        raise NoEncontrado(f"Reporte desconocido: {nombre}")
    return REPORTES[nombre]


@lru_cache
def lote(nombre: str) -> ReporteLote | None:
    """El `ReporteLote` del módulo, o None si el reporte tiene lógica propia (cmg-mora, bancarizados…)."""
    valor = getattr(import_module(reporte(nombre).modulo), "REPORTE", None)
    return valor if isinstance(valor, ReporteLote) else None


def _resumen(r: Reporte) -> dict:
    lt = lote(r.nombre)
    return dict(
        nombre=r.nombre, grupo=r.grupo, orden=r.orden, frecuencia=r.frecuencia, servidores=list(r.servidores),
        descripcion=r.descripcion, carpeta=r.carpeta, es_lote=lt is not None,
        escribe_en_bd=bool(lt and lt.escribe_en_bd), vacio_valido=bool(lt and lt.vacio_valido),
        envia_correo=bool(lt and lt.entrega is not None), avisos=list(lt.avisos) if lt else [],
    )


def catalogo() -> list[e.ReporteResumen]:
    return [e.ReporteResumen(**_resumen(r)) for r in ordenados()]


def corte_por_defecto(nombre: str) -> e.Corte:
    r = reporte(nombre)
    try:
        fecha, origen = resolver_corte(r.frecuencia, None, lunes_sabado=nombre not in DIA_ANTERIOR_SIMPLE)
        return e.Corte(fecha=fecha, origen=origen)
    except ConfiguracionError as exc:
        return e.Corte(fecha=None, origen=str(exc))


def detalle(nombre: str) -> e.ReporteDetalle:
    r = reporte(nombre)
    corte = corte_por_defecto(nombre)
    ref = corte.fecha or date.today()
    tablas = [
        e.TablaReporte(nombre=nombre_resuelto(t, ref), servidor=t.servidor, tipo=t.tipo, columna_fecha=t.col_fecha,
                       confianza=t.confianza, verificable=t.verificable, condicion=regla_de(nombre, t.nombre))
        for t in tablas_de(nombre)
    ]
    return e.ReporteDetalle(**_resumen(r), corte=corte, tablas=tablas)


# ---------------------------------------------------------------- validación de la fecha de corte

def validar_corte(nombre: str, fecha: date | None) -> date:
    """Misma regla que el ejecutor: --fecha-corte > .env > (diaria) día anterior; mensual = fin de mes."""
    r = reporte(nombre)
    try:
        corte, _ = resolver_corte(r.frecuencia, fecha, lunes_sabado=nombre not in DIA_ANTERIOR_SIMPLE)
        if r.frecuencia == "mensual":
            Cortes.mensual(corte)
    except ConfiguracionError as exc:
        raise PeticionInvalida(str(exc)) from exc
    if corte > date.today():
        raise PeticionInvalida(f"La fecha de corte {corte:%Y-%m-%d} está en el futuro.")
    return corte


# ---------------------------------------------------------------- verificación y solicitud a Producción

def verificar(nombre: str, fecha: date | None) -> e.Verificacion:
    corte = validar_corte(nombre, fecha)
    resultados = verificar_reporte(nombre, corte)
    pendientes = [x for x in resultados if x.estado in {Estado.DESACTUALIZADA, Estado.NO_EXISTE}]
    return e.Verificacion(
        corte=corte,
        listo=not pendientes and not any(x.estado is Estado.ERROR for x in resultados),
        resultados=[
            e.ResultadoTabla(nombre=nombre_resuelto(x.tabla, corte), servidor=x.tabla.servidor, estado=x.estado.value,
                             ultima_fecha=x.ultima_fecha, detalle=x.detalle)
            for x in resultados
        ],
        solicitud=mensaje_solicitud(nombre, corte, resultados) if pendientes else None,
    )


def guardar_solicitud(nombre: str, fecha: date | None) -> e.Solicitud:
    """Verifica de nuevo (no se confía en lo que mande el navegador) y guarda el mensaje en data/outputs/solicitudes."""
    corte = validar_corte(nombre, fecha)
    texto = mensaje_solicitud(nombre, corte, verificar_reporte(nombre, corte))
    destino = DIR_OUTPUTS / "solicitudes"
    destino.mkdir(parents=True, exist_ok=True)
    ruta = destino / f"solicitud_{nombre}_{corte:%Y%m%d}.txt"
    ruta.write_text(texto + "\n", encoding="utf-8")
    return e.Solicitud(archivo=ruta.name, texto=texto)


def _corte(frecuencia: str) -> e.Corte:
    try:
        fecha, origen = resolver_corte(frecuencia, None)
        return e.Corte(fecha=fecha, origen=origen)
    except ConfiguracionError as exc:
        return e.Corte(fecha=None, origen=str(exc))


def configuracion() -> e.ConfiguracionGeneral:
    return e.ConfiguracionGeneral(corte_mensual=_corte("mensual"), corte_diario=_corte("diaria"),
                                  dir_inputs=str(DIR_INPUTS), dir_outputs=str(DIR_OUTPUTS), driver_odbc=DRIVER_ODBC)


def servidores() -> list[e.Servidor]:
    return [
        e.Servidor(nombre=s.nombre, servidor=s.servidor, autenticacion="SQL" if (s.usuario or not s.windows_auth_defecto) else "Windows",
                   credenciales_en_env=bool(s.usuario and s.clave), descripcion=s.descripcion, bases=list(BASES_DE[s.nombre]))
        for s in SERVIDORES.values()
    ]


def probar_servidor(nombre: str) -> e.PruebaConexion:
    """SELECT 1 contra un servidor (mismo camino que usan los reportes: db.leer_sql)."""
    if nombre not in SERVIDORES:
        raise NoEncontrado(f"Servidor desconocido: {nombre}")
    inicio = time.perf_counter()
    try:
        leer_sql(nombre, "SELECT 1 AS ok")
        ok, detalle = True, "Conexión correcta"
    except Exception as exc:  # noqa: BLE001 - diagnóstico para la pantalla, igual que probar_conexiones
        ok, detalle = False, f"{type(exc).__name__}: {str(exc)[:200]}"
    return e.PruebaConexion(nombre=nombre, ok=ok, detalle=detalle, milisegundos=round((time.perf_counter() - inicio) * 1000))


# ---------------------------------------------------------------- argumentos de ejecución

def argumentos(p: e.PedidoEjecucion) -> list[str]:
    """Valida el pedido con las reglas del reporte y devuelve los argumentos de `main.py`."""
    nombre = p.reporte
    corte = validar_corte(nombre, p.fecha_corte)
    lt = lote(nombre)
    if lt is None:
        if p.forzar or p.confirmar_escritura or p.correo == "todos":
            raise PeticionInvalida(f"«{nombre}» tiene lógica propia: no admite forzar, confirmar escritura ni envío de correo.")
        return [nombre, "--mes", f"{corte:%Y-%m}"] if nombre in USAN_MES else [nombre, "--fecha-corte", f"{corte:%Y-%m-%d}"]
    if lt.escribe_en_bd and not p.confirmar_escritura:
        raise PeticionInvalida(f"«{nombre}» crea/borra tablas permanentes: confirma la escritura en la base de datos.")
    args = [nombre, "--fecha-corte", f"{corte:%Y-%m-%d}"]
    if p.forzar:
        args.append("--forzar")
    if p.confirmar_escritura:
        args.append("--confirmar-escritura")
    if lt.entrega is not None:
        if p.correo == "todos":
            if not p.conforme:
                raise PeticionInvalida("Para enviar a toda la lista debes confirmar que revisaste el correo de prueba.")
            envio = estado_envio(nombre, corte)
            if not envio.prueba:
                raise PeticionInvalida(f"No hay correo de prueba para {corte:%Y-%m-%d}: primero ejecuta el reporte (envía la prueba).")
            if envio.todos:
                raise PeticionInvalida(f"Ya se envió a toda la lista el {envio.todos}.")
            args.append("--conforme")
        args += ["--correo", p.correo]
    elif p.correo == "todos":
        raise PeticionInvalida(f"«{nombre}» no envía correo.")
    return args


# ---------------------------------------------------------------- archivos generados y vista previa

def _carpeta(nombre: str) -> Path:
    reporte(nombre)
    return carpeta_salida(nombre)


def _ruta_segura(nombre: str, archivo: str) -> Path:
    """Solo archivos directamente dentro de la carpeta de salida del reporte (sin rutas relativas)."""
    carpeta = _carpeta(nombre).resolve()
    ruta = (carpeta / archivo).resolve()
    if ruta.parent != carpeta or not ruta.is_file():
        raise NoEncontrado(f"No existe el archivo {archivo}")
    return ruta


def archivos(nombre: str) -> list[e.Archivo]:
    carpeta = _carpeta(nombre)
    if not carpeta.is_dir():
        return []
    salida = [
        e.Archivo(nombre=f.name, tipo=TIPOS.get(f.suffix.lower(), "otro"), tamano=f.stat().st_size,
                  modificado=datetime.fromtimestamp(f.stat().st_mtime).isoformat(timespec="seconds"))
        for f in carpeta.iterdir() if f.is_file() and not f.name.startswith("~$")
    ]
    return sorted(salida, key=lambda a: a.modificado, reverse=True)


def ruta_descarga(nombre: str, archivo: str) -> Path:
    return _ruta_segura(nombre, archivo)


def _celda(v):
    if v is None or isinstance(v, (int, float, str)):
        return v
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v)


def vista_previa(nombre: str, archivo: str) -> e.VistaPrevia:
    ruta = _ruta_segura(nombre, archivo)
    tipo = TIPOS.get(ruta.suffix.lower())
    if tipo == "texto":
        return e.VistaPrevia(archivo=ruta.name, tipo="texto", texto=ruta.read_text(encoding="utf-8", errors="replace")[:MAX_TEXTO])
    if tipo != "excel":
        raise PeticionInvalida("Solo hay vista previa de Excel y texto; las imágenes se abren directamente.")
    from openpyxl import load_workbook

    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hojas = []
        for ws in libro.worksheets:
            filas = ws.iter_rows(values_only=True)
            columnas = [str(c) if c is not None else "" for c in next(filas, ())]
            datos = [[_celda(v) for v in fila] for _, fila in zip(range(FILAS_VISTA_PREVIA), filas)]
            hojas.append(e.HojaPrevia(nombre=ws.title, columnas=columnas, filas=datos, total_filas=max((ws.max_row or 1) - 1, 0)))
        return e.VistaPrevia(archivo=ruta.name, tipo="excel", hojas=hojas)
    finally:
        libro.close()


# ---------------------------------------------------------------- envío por correo (prueba -> conforme -> todos)

def estado_envio(nombre: str, fecha: date | None) -> e.EstadoEnvio:
    lt = lote(nombre)
    if lt is None or lt.entrega is None:
        raise NoEncontrado(f"«{nombre}» no envía correo.")
    corte = validar_corte(nombre, fecha)
    ruta = _carpeta(nombre) / f"estado_envio_{corte:%Y%m%d}.json"  # mismo archivo que escribe EntregaCorreoResumen
    datos = json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else {}
    excel = datos.get("excel")
    return e.EstadoEnvio(corte=corte, prueba=datos.get("prueba"), todos=datos.get("todos"), excel=Path(excel).name if excel else None)

"""Lectura de archivos generados para la vista previa (Excel con openpyxl)."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from modulos.reportes.aplicacion import dto as dto_reportes



def _celda(v):
    if v is None or isinstance(v, (int, float, str)):
        return v
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v)


def hojas_excel(ruta: Path, max_filas: int) -> list[dto_reportes.HojaPrevia]:
    """Cada hoja con sus columnas y las primeras `max_filas` filas (solo lectura, valores calculados)."""
    from openpyxl import load_workbook

    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hojas = []
        for ws in libro.worksheets:
            filas = ws.iter_rows(values_only=True)
            columnas = [str(c) if c is not None else "" for c in next(filas, ())]
            datos = [[_celda(v) for v in fila] for _, fila in zip(range(max_filas), filas)]
            hojas.append(dto_reportes.HojaPrevia(nombre=ws.title, columnas=columnas, filas=datos, total_filas=max((ws.max_row or 1) - 1, 0)))
        return hojas
    finally:
        libro.close()

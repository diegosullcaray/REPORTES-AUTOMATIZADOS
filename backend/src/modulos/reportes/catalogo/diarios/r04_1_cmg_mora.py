"""CMG Mora (diario): recaudo + provisiones -> INSERTs para [storage].[com_act].[SBTVRIE001].

Conexión: servidor rcc (172.20.0.70), base DW_Raw_v2 (dbriesgos por nombre de 3 partes). Ver core/configuracion.py / .env.

Las operaciones diarias (recaudo + gasto) cargan el mismo día; las provisiones (PROV_PROY) unos 5 días después. Por eso hay fases:
  --fase operaciones  -> operaciones_<fecha>.xlsx (recaudo + gasto)
  --fase provisiones  -> provisiones_<fecha>.xlsx (SSTKPROV; solo si la tabla PROV_PROY ya existe)
  --fase generar      -> con los dos Excel del día arma los INSERT (TXT)
Sin --fase hace todo de una vez (las dos tablas ya cargadas).
"""
import datetime
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook

from modulos.reportes.infraestructura.entregas import chat as entrega_chat
from modulos.reportes.dominio.fechas import fecha_iso
from modulos.reportes.aplicacion.cortes import resolver_corte
from core.excepciones import ConfiguracionError
from modulos.reportes.aplicacion.rutas import carpeta_salida
from core.bd import conexion_pyodbc

TABLA = "DW_Raw_v2.dbo.CMGMora_Recaudo"

# Se incluye filtro por fecha para traer solo la data del día objetivo
SQL_RECAUDO = f"""
TRUNCATE TABLE {TABLA};

;WITH recaudo as(
    SELECT
        CAST(FECHA_CIERRE as date) FECHA_CIERRE
        ,CAST(SALDO as float) SALDO
        ,CAST(RECAUDO as float) RECAUDO
        ,CAST(TASA as float) TASA
        ,CORREDOR ,GRUPO ,TERRITORIO
        ,LTRIM(RTRIM(TERRITORIO))+'-'+LTRIM(RTRIM(GRUPO))+'-'+LTRIM(RTRIM(CORREDOR)) TER_GRU_COR
        ,CASE WHEN CORREDOR IN ('SIN ASIGNAR', 'TOTAL') THEN 20 ELSE 19 END TIP_COD
    FROM dbriesgos.dbo.RECAUDO_DIARIO_FINANZAS
    WHERE CAST(FECHA_CIERRE as date) = ?
)
INSERT INTO {TABLA}
SELECT
    B.FECHA_CIERRE AS SFECPRO
    ,B.CORREDOR ,B.GRUPO ,B.TERRITORIO
    ,B.TIP_COD AS STIPCOD
    ,J.RCODCOR AS SCODREL
    ,B.RECAUDO AS SMONREC
    ,B.TASA AS STASREC
    ,0 AS SRECCAST12M
    ,0 AS SSTKPROV
    ,0 AS SGASPROVCART
    ,0 AS SGASPROVCONTA
    ,0 AS SGASPROVOM3155
    ,0 AS SGASPROVVOLU
    ,0 AS SGASPROVBRUTO
    ,0 AS SGASPROVBRUTO12M
FROM recaudo B
LEFT JOIN DW_Raw_v2.dbo.CMGMora_STRJERCOR J ON (B.TER_GRU_COR=J.[TER-GRU-COR]);

UPDATE {TABLA}
SET STIPCOD = 7
WHERE CORREDOR IN ('TOTAL') AND GRUPO IN ('TOTAL');
"""

# Asignamos variables estáticas
SQL_CONSTANTES = f"""
UPDATE {TABLA}
SET
    SRECCAST12M = 3437395.96
    ,SGASPROVBRUTO12M = 69127536.6
    ,SGASPROVOM3155 = 0
WHERE STIPCOD = 7;
"""

SQL_INSERTS = f"""
SELECT
    'INSERT INTO [storage].[com_act].[SBTVRIE001] VALUES (' +
    '''' + CONVERT(VARCHAR(25), SFECPRO, 121) + ''',' +
    CAST(STIPCOD AS VARCHAR(50)) + ', ' +
    '''' + REPLACE(SCODREL, '''', '''''') + ''', ' +
    CAST(CAST(SMONREC AS DECIMAL(20,7)) AS VARCHAR(50))  + ', ' +
    CAST(CAST(STASREC AS DECIMAL(20,7)) AS VARCHAR(50))  + ', ' +
    CAST(SRECCAST12M AS VARCHAR(50)) + ', ' +
    CAST(SSTKPROV AS VARCHAR(50)) + ', ' +
    CAST(SGASPROVCART AS VARCHAR(50)) + ', ' +
    CAST(SGASPROVCONTA AS VARCHAR(50)) + ', ' +
    CAST(SGASPROVOM3155 AS VARCHAR(50)) + ', ' +
    CAST(SGASPROVVOLU AS VARCHAR(50)) + ', ' +
    CAST(SGASPROVBRUTO AS VARCHAR(50)) + ', ' +
    CAST(SGASPROVBRUTO12M AS VARCHAR(50)) +
    ');' AS Inserts
FROM {TABLA};
"""


def _cargar_recaudo(cursor, fec_str: str) -> bool:
    """PARTE 1: carga y tipifica el recaudo del día. False = no hay recaudo para esa fecha."""
    print("-> Ejecutando PARTE 1 (Carga y tipificación en DW_Raw_v2)...")
    cursor.execute(SQL_RECAUDO, fec_str)
    cursor.execute(f"SELECT COUNT(*) FROM {TABLA}")
    if cursor.fetchone()[0] == 0:
        print(f"\n[!] ALERTA: No se encontró data en RECAUDO_DIARIO_FINANZAS para la fecha {fec_str}. Abortando.")
        return False
    return True


def _gasto(cursor, fec_str: str) -> float:
    cursor.execute("SELECT ISNULL(SUM(GASTO_OPE), 0.0) FROM [dbriesgos].[dbo].GASTO_PROV_OPE_DIARIA WHERE CAST(FC_DIA as date) = ?", fec_str)
    return float(cursor.fetchone()[0] or 0.0)


def _stkprov(cursor, tabla_prov: str) -> float | None:
    """Suma de la provisión proyectada; None si la tabla del día aún no existe."""
    cursor.execute(f"SELECT OBJECT_ID('{tabla_prov}')")
    if cursor.fetchone()[0] is None:
        print(f"\n[!] ERROR: La tabla {tabla_prov} no existe en la base de datos. Abortando.")
        return None
    cursor.execute(f"SELECT ISNULL(SUM(PROVISION_PROYECTADA), 0.0) FROM {tabla_prov}")
    return float(cursor.fetchone()[0] or 0.0)


def _en_cero(**variables: float) -> bool:
    """VALIDACIÓN CLAVE: variable crítica en 0 => no se generan los INSERT."""
    print("   Variables -> " + " | ".join(f"{k}: {v:.2f}" for k, v in variables.items()))
    if all(variables.values()):
        return False
    print("\n" + "=" * 60)
    print("🛑 ADVERTENCIA: Las variables de provisión están en 0.")
    print("🛑 SEGÚN LAS REGLAS, NO SE GENERARÁN LOS INSERTS. ABORTANDO.")
    print("=" * 60 + "\n")
    return True


def _poner_gasto(cursor, gasto: float) -> None:
    cursor.execute(f"UPDATE {TABLA} SET SGASPROVCART = ?, SGASPROVCONTA = ? WHERE STIPCOD = 7", gasto, gasto)  # contable = cartera


def _poner_stkprov(cursor, stkprov: float) -> None:
    cursor.execute(f"UPDATE {TABLA} SET SSTKPROV = ? WHERE STIPCOD = 7", stkprov)


def _escribir_inserts(cursor, ruta_salida: str, fec_str: str) -> int:
    """PARTE 3: sentencias INSERT finales al TXT del día."""
    print("-> Extrayendo sentencias INSERT finales (PARTE 3)...")
    cursor.execute(SQL_INSERTS)
    inserts = cursor.fetchall()
    if not inserts:
        print("\n[!] No se generaron registros en la tabla.")
        return 1
    archivo_txt = os.path.join(ruta_salida, f"inserts_{fec_str}.txt")
    with open(archivo_txt, 'w', encoding='utf-8') as f:
        for fila in inserts:
            f.write(fila[0] + "\n")
    print(f"\n✅ [ÉXITO] Se generaron {len(inserts)} registros INSERTS.")
    print(f"✅ [ARCHIVO GUARDADO EN] {archivo_txt}")
    return 0


def _guardar_excel(ruta: str, cabecera: list, filas: list) -> None:
    wb = Workbook()
    wb.active.append(cabecera)
    for fila in filas:
        wb.active.append(list(fila))
    wb.save(ruta)
    print(f"\n✅ [ÉXITO] Archivo guardado: {ruta}")


def _leer_excel(ruta: str) -> tuple[list, list]:
    if not os.path.exists(ruta):
        raise FileNotFoundError(f"Falta {os.path.basename(ruta)}: genera esa fase primero.")
    cabecera, *filas = load_workbook(ruta, read_only=True, data_only=True).active.iter_rows(values_only=True)
    return list(cabecera), filas


def generar_inserts_sql(fecha_corte: datetime.date | None = None, fase: str | None = None) -> int:
    """0 = archivo de la fase generado. 1 = abortó o falló (sin datos, tabla del día ausente, provisiones en 0, error SQL): nunca se presenta como éxito."""
    # Fecha: --fecha-corte > FECHA_CORTE_DIARIA del .env > día anterior (lunes: sábado)
    hoy = datetime.date.today()
    fecha_ejec, origen = resolver_corte("diaria", fecha_corte)
    print(f"Fecha actual: {hoy.strftime('%Y-%m-%d')}")
    print(f"Fecha de ejecución objetivo: {fecha_ejec.strftime('%Y-%m-%d')} ({origen})")

    fec_str = fecha_ejec.strftime('%Y-%m-%d')
    tabla_prov = f"[dbriesgos].[dbo].[PROV_PROY_{fecha_ejec.strftime('%Y%m%d')}_0]"
    ruta_salida = carpeta_salida('cmg-mora')
    os.makedirs(ruta_salida, exist_ok=True)
    xlsx_oper = os.path.join(ruta_salida, f"operaciones_{fec_str}.xlsx")
    xlsx_prov = os.path.join(ruta_salida, f"provisiones_{fec_str}.xlsx")

    try:
        print("\nConectando a la base de datos SQL Server...")
        conn_ctx = conexion_pyodbc('rcc', base='DW_Raw_v2')
        conn = conn_ctx.__enter__()
        cursor = conn.cursor()

        if fase == "generar":
            # Los dos Excel del día vuelven a la tabla de trabajo y de ahí salen los INSERT con el mismo formato de siempre
            cab, filas = _leer_excel(xlsx_oper)
            (cab_prov, filas_prov) = _leer_excel(xlsx_prov)
            stkprov = float(filas_prov[0][cab_prov.index("SSTKPROV")] or 0.0)
            cursor.execute(f"TRUNCATE TABLE {TABLA}")
            cursor.executemany(f"INSERT INTO {TABLA} ({', '.join(cab)}) VALUES ({', '.join('?' * len(cab))})", filas)
            if _en_cero(SSTKPROV=stkprov):
                return 1
            _poner_stkprov(cursor, stkprov)
            return _escribir_inserts(cursor, ruta_salida, fec_str)

        if fase == "provisiones":
            stkprov = _stkprov(cursor, tabla_prov)
            if stkprov is None or _en_cero(SSTKPROV=stkprov):
                return 1
            _guardar_excel(xlsx_prov, ["SSTKPROV"], [(stkprov,)])
            return 0

        if not _cargar_recaudo(cursor, fec_str):
            return 1
        print("-> Ejecutando PARTE 2 (Cálculo de provisiones)...")
        cursor.execute(SQL_CONSTANTES)
        gasto = _gasto(cursor, fec_str)

        if fase == "operaciones":
            if _en_cero(SGASPROVCART=gasto):
                return 1
            _poner_gasto(cursor, gasto)
            cursor.execute(f"SELECT * FROM {TABLA}")
            _guardar_excel(xlsx_oper, [c[0] for c in cursor.description], cursor.fetchall())
            return 0

        stkprov = _stkprov(cursor, tabla_prov)
        if stkprov is None or _en_cero(SSTKPROV=stkprov, SGASPROVCART=gasto):
            return 1
        _poner_gasto(cursor, gasto)
        _poner_stkprov(cursor, stkprov)
        return _escribir_inserts(cursor, ruta_salida, fec_str)

    except ConfiguracionError:
        raise  # lo traduce main() a código 2 (configuración), no a un error de datos
    except Exception as e:
        print(f"\n[X] Ocurrió un error en la ejecución SQL o en el proceso: {e}")
        return 1
    finally:
        if 'conn_ctx' in locals():
            conn_ctx.__exit__(None, None, None)
            print("-> Conexión a la base de datos cerrada de manera segura.")

def enviar_por_chat(fecha_corte: datetime.date | None) -> int:
    """Envía el TXT del día por Google Chat a los destinatarios de ENTREGA_CHAT_CMG_MORA. 0 = enviado; 1 = no se pudo."""
    fecha_ejec, _ = resolver_corte("diaria", fecha_corte)
    archivo = Path(carpeta_salida("cmg-mora")) / f"inserts_{fecha_ejec:%Y-%m-%d}.txt"
    texto = f"*CMG Mora* · corte {fecha_ejec:%d/%m/%Y}\nAdjunto los INSERT de SBTVRIE001 ({archivo.name})."
    if entrega_chat.enviar("cmg-mora", [archivo], texto):
        return 0
    print("✗ El TXT se generó, pero no se envió por Google Chat (descárgalo desde la pestaña Archivos o reintenta el envío).")
    return 1


def main(argv=None) -> int:
    import argparse

    parser = argparse.ArgumentParser(prog="main.py cmg-mora", description="CMG Mora: recaudo + provisiones -> INSERTs.")
    parser.add_argument("--fecha-corte", type=lambda v: fecha_iso(v), help="AAAA-MM-DD (defecto: FECHA_CORTE_DIARIA del .env; sin ninguna: día anterior, lunes = sábado)")
    parser.add_argument("--fase", choices=["operaciones", "provisiones", "generar"], help="una sola fase del flujo (defecto: todo de una vez)")
    parser.add_argument("--chat", action="store_true", help="tras generar el TXT, enviarlo por Google Chat a ENTREGA_CHAT_CMG_MORA del .env")
    args = parser.parse_args(argv)
    if args.chat and args.fase in ("operaciones", "provisiones"):
        parser.error("--chat solo aplica cuando se generan los INSERT")
    try:
        rc = generar_inserts_sql(args.fecha_corte, args.fase)
        return enviar_por_chat(args.fecha_corte) if rc == 0 and args.chat else rc
    except (ConfiguracionError, ValueError) as exc:
        print(f"✗ {exc}")
        return 2


if __name__ == '__main__':
    raise SystemExit(main())

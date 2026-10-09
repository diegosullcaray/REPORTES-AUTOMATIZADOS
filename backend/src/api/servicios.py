"""Casos de uso de la interfaz web sobre el motor de `reportes` (catálogo, tablas, verificación, solicitud, archivos, envío).

No reimplementa reglas: llama a las mismas funciones que usa la consola, así la web y `main.py` no pueden divergir.
Las validaciones lanzan `PeticionInvalida` (422) o `NoEncontrado` (404); la capa HTTP solo las traduce.
"""

from __future__ import annotations

import getpass
import json
import os
import platform
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from decimal import Decimal
from functools import lru_cache
from importlib import import_module
from pathlib import Path

from reportes.comun import correo
from reportes.comun.ejecutor import ReporteLote
from reportes.comun.fechas import Cortes, fecha_iso, resolver_corte
from dotenv import dotenv_values

from reportes.config import BASES_DE, DIR_INPUTS, DIR_OUTPUTS, DRIVER_ODBC, RAIZ, SERVIDORES, ConfiguracionError
from reportes.db import leer_sql
from reportes.registro import DIA_ANTERIOR_SIMPLE, REPORTES, Reporte, carpeta_salida, ordenados
from reportes.reglas_fecha import regla_de
from reportes.tablas import tablas_de
from reportes.verificacion import Estado, mensaje_solicitud, mensaje_solicitud_grupo, nombre_resuelto, Resultado, verificar_reporte, verificar_tabla

from . import entorno, sesion
from . import esquemas as e

FILAS_VISTA_PREVIA = 200
MAX_TEXTO = 200_000
# Reportes de Erick que reciben el mes (AAAA-MM) en lugar de la fecha de corte.
USAN_MES = frozenset({"bancarizados-producto", "indicadores-clientes"})
TIPOS = {".xlsx": "excel", ".xlsm": "excel", ".jpg": "imagen", ".jpeg": "imagen", ".png": "imagen", ".txt": "texto", ".csv": "texto", ".json": "texto"}


class PeticionInvalida(ValueError):
    """La petición no cumple una regla del reporte (fecha, confirmación, correo…)."""


class NoEncontrado(LookupError):
    pass


# ---------------------------------------------------------------- catálogo

def reporte(nombre: str) -> Reporte:
    if nombre not in REPORTES:
        raise NoEncontrado(f"Reporte desconocido: {nombre}")
    return REPORTES[nombre]


@lru_cache
def lote(nombre: str) -> ReporteLote | None:
    """El `ReporteLote` del módulo, o None si el reporte tiene lógica propia (cmg-mora, bancarizados…)."""
    valor = getattr(import_module(reporte(nombre).modulo), "REPORTE", None)
    return valor if isinstance(valor, ReporteLote) else None


# Reportes con lógica propia (sin ReporteLote): lo que la pantalla debe advertir antes de ejecutar.
AVISOS_PROPIOS = {
    "cmg-mora": [
        "Trunca y recarga DW_Raw_v2.CMGMora_Recaudo (rcc) y genera los INSERT de SBTVRIE001. Termina en error si no hay recaudo del día, "
        "falta la tabla PROV_PROY del día o las provisiones están en 0. No lo ejecutes a la vez desde la consola.",
    ],
}


def _resumen(r: Reporte) -> dict:
    lt = lote(r.nombre)
    return dict(
        nombre=r.nombre, grupo=r.grupo, orden=r.orden, frecuencia=r.frecuencia, servidores=list(r.servidores),
        descripcion=r.descripcion, carpeta=r.carpeta, es_lote=lt is not None,
        escribe_en_bd=bool(lt and lt.escribe_en_bd), vacio_valido=bool(lt and lt.vacio_valido),
        envia_correo=bool(lt and lt.entrega is not None), avisos=list(lt.avisos) if lt else AVISOS_PROPIOS.get(r.nombre, []),
    )


def catalogo() -> list[e.ReporteResumen]:
    return [e.ReporteResumen(**_resumen(r)) for r in ordenados()]


def corte_por_defecto(nombre: str) -> e.Corte:
    r = reporte(nombre)
    try:
        fecha, origen = resolver_corte(r.frecuencia, None, lunes_sabado=nombre not in DIA_ANTERIOR_SIMPLE)
        return e.Corte(fecha=fecha, origen=origen)
    except ConfiguracionError as exc:
        return e.Corte(fecha=None, origen=str(exc))


def detalle(nombre: str) -> e.ReporteDetalle:
    r = reporte(nombre)
    corte = corte_por_defecto(nombre)
    ref = corte.fecha or date.today()
    tablas = [
        e.TablaReporte(nombre=nombre_resuelto(t, ref), servidor=t.servidor, tipo=t.tipo, columna_fecha=t.col_fecha,
                       confianza=t.confianza, verificable=t.verificable, condicion=regla_de(nombre, t.nombre))
        for t in tablas_de(nombre)
    ]
    return e.ReporteDetalle(**_resumen(r), corte=corte, tablas=tablas)


# ---------------------------------------------------------------- validación de la fecha de corte

def validar_corte(nombre: str, fecha: date | None) -> date:
    """Misma regla que el ejecutor: --fecha-corte > .env > (diaria) día anterior; mensual = fin de mes."""
    r = reporte(nombre)
    try:
        corte, _ = resolver_corte(r.frecuencia, fecha, lunes_sabado=nombre not in DIA_ANTERIOR_SIMPLE)
        if r.frecuencia == "mensual":
            Cortes.mensual(corte)
    except ConfiguracionError as exc:
        raise PeticionInvalida(str(exc)) from exc
    if corte > date.today():
        raise PeticionInvalida(f"La fecha de corte {corte:%Y-%m-%d} está en el futuro.")
    return corte


# ---------------------------------------------------------------- verificación y solicitud a Producción

def verificar(nombre: str, fecha: date | None) -> e.Verificacion:
    corte = validar_corte(nombre, fecha)
    resultados = verificar_reporte(nombre, corte)
    pendientes = [x for x in resultados if x.estado in {Estado.DESACTUALIZADA, Estado.NO_EXISTE}]
    return e.Verificacion(
        corte=corte,
        listo=not pendientes and not any(x.estado is Estado.ERROR for x in resultados),
        resultados=[
            e.ResultadoTabla(nombre=nombre_resuelto(x.tabla, corte), servidor=x.tabla.servidor, estado=x.estado.value,
                             ultima_fecha=x.ultima_fecha, detalle=x.detalle)
            for x in resultados
        ],
        solicitud=mensaje_solicitud(nombre, corte, resultados) if pendientes else None,
    )


GRUPOS_VALIDACION = {"piero": "Heredados de Piero", "erick": "Heredados de Erick"}
_PENDIENTE = {Estado.DESACTUALIZADA, Estado.NO_EXISTE}


def _verificar_grupo(grupo: str, fecha: date | None):
    """Verifica cada tabla distinta una sola vez aunque la usen varios reportes. Devuelve corte, reportes y {tabla: [Tabla, Resultado, reportes]}."""
    if grupo not in GRUPOS_VALIDACION:
        raise NoEncontrado(f"Grupo desconocido: {grupo}")
    reportes = [r for r in ordenados() if r.grupo == grupo and r.frecuencia == "mensual" and tablas_de(r.nombre)]
    if not reportes:
        raise NoEncontrado(f"El grupo «{grupo}» no tiene reportes mensuales con tablas registradas.")
    corte = validar_corte(reportes[0].nombre, fecha)
    unicas: dict[str, list] = {}
    for r in reportes:
        for t in tablas_de(r.nombre):
            clave = nombre_resuelto(t, corte)
            if clave not in unicas:
                unicas[clave] = [t, None, []]
            unicas[clave][2].append(r.nombre)

    def por_servidor(claves: list[str]) -> None:
        """Un servidor, en orden. Si no responde, sus demás tablas se marcan con la misma causa sin reintentar (cada intento tarda)."""
        caido: Resultado | None = None
        for clave in claves:
            t = unicas[clave][0]
            if caido and t.verificable:
                unicas[clave][1] = Resultado(t, Estado.ERROR, detalle=caido.detalle, sin_conexion=True)
            else:
                unicas[clave][1] = verificar_tabla(t, corte)
                if unicas[clave][1].sin_conexion:
                    caido = unicas[clave][1]

    grupos: dict[str, list[str]] = {}
    for clave, (t, _, _) in unicas.items():
        grupos.setdefault(t.servidor, []).append(clave)
    with ThreadPoolExecutor(max_workers=len(grupos)) as pool:  # los servidores se consultan a la vez
        list(pool.map(por_servidor, grupos.values()))
    return corte, reportes, unicas


def _solucion(causa: str) -> str:
    if "driver ODBC" in causa:
        return "Instala «ODBC Driver 17 for SQL Server» o define DB_ODBC_DRIVER en backend/.env con un driver instalado."
    if "requiere" in causa and "_USER" in causa:
        return "Define el usuario y la contraseña de esa conexión en backend/.env (por ejemplo RCC_USER y RCC_PASSWORD)."
    if "No se pudo conectar" in causa:
        return "Comprueba la red o la VPN y que el nombre del servidor en backend/.env sea correcto; luego vuelve a verificar."
    if "rechazó el inicio de sesión" in causa:
        return "Revisa el usuario y la contraseña de esa conexión en backend/.env."
    return "Revisa el detalle de la tabla y vuelve a verificar."


def verificar_grupo(grupo: str, fecha: date | None) -> e.VerificacionGrupo:
    corte, reportes, unicas = _verificar_grupo(grupo, fecha)
    resumen = []
    for r in reportes:
        propias = [x[1] for x in unicas.values() if r.nombre in x[2]]
        pendientes = sum(x.estado in _PENDIENTE for x in propias)
        dudosas = sum(x.estado is Estado.ERROR for x in propias)
        resumen.append(e.ResumenValidacion(nombre=r.nombre, orden=r.orden, tablas=len(propias), pendientes=pendientes, dudosas=dudosas, listo=not (pendientes or dudosas)))
    por_pedir = [tuple(x) for x in unicas.values() if x[1].estado in _PENDIENTE]
    causas: dict[str, int] = {}
    for _, res, _ in unicas.values():
        if res.estado is Estado.ERROR:
            causas[res.detalle] = causas.get(res.detalle, 0) + 1
    return e.VerificacionGrupo(
        problemas=[e.ProblemaValidacion(causa=c, tablas=n, solucion=_solucion(c)) for c, n in sorted(causas.items(), key=lambda x: -x[1])],
        grupo=grupo, titulo=GRUPOS_VALIDACION[grupo], corte=corte, listo=all(x.listo for x in resumen), reportes=resumen,
        tablas=[
            e.TablaValidada(nombre=nombre, servidor=t.servidor, estado=res.estado.value, ultima_fecha=res.ultima_fecha, detalle=res.detalle, reportes=usan)
            for nombre, (t, res, usan) in unicas.items()
        ],
        solicitud=mensaje_solicitud_grupo(GRUPOS_VALIDACION[grupo], corte, por_pedir) if por_pedir else None,
    )


def guardar_solicitud_grupo(grupo: str, fecha: date | None) -> e.Solicitud:
    """Vuelve a verificar (no se confía en el navegador) y guarda el mensaje único en data/outputs/solicitudes."""
    corte, _, unicas = _verificar_grupo(grupo, fecha)
    texto = mensaje_solicitud_grupo(GRUPOS_VALIDACION[grupo], corte, [tuple(x) for x in unicas.values() if x[1].estado in _PENDIENTE])
    destino = DIR_OUTPUTS / "solicitudes"
    destino.mkdir(parents=True, exist_ok=True)
    ruta = destino / f"solicitud_mensual_{grupo}_{corte:%Y%m%d}.txt"
    ruta.write_text(texto + "\n", encoding="utf-8")
    return e.Solicitud(archivo=ruta.name, texto=texto)


def guardar_solicitud(nombre: str, fecha: date | None) -> e.Solicitud:
    """Verifica de nuevo (no se confía en lo que mande el navegador) y guarda el mensaje en data/outputs/solicitudes."""
    corte = validar_corte(nombre, fecha)
    texto = mensaje_solicitud(nombre, corte, verificar_reporte(nombre, corte))
    destino = DIR_OUTPUTS / "solicitudes"
    destino.mkdir(parents=True, exist_ok=True)
    ruta = destino / f"solicitud_{nombre}_{corte:%Y%m%d}.txt"
    ruta.write_text(texto + "\n", encoding="utf-8")
    return e.Solicitud(archivo=ruta.name, texto=texto)


def perfil() -> e.Perfil:
    try:
        destinatarios: int | None = len(correo.leer_destinatarios())
    except (OSError, ConfiguracionError, correo.CorreoError):
        destinatarios = None
    return e.Perfil(
        usuario=getpass.getuser(), equipo=platform.node(),
        correo_prueba=os.getenv("CORREO_PRUEBA", "diego.sullcaray@confianza.pe").strip(),  # mismo defecto que correo.config_desde_env
        cuenta_envio=os.getenv("SMTP_USER", "").strip() or None,
        clave_envio_configurada=bool(os.getenv("SMTP_PASSWORD", "").strip()),
        webhook_configurado=bool(os.getenv("GOOGLE_CHAT_WEBHOOK_URL", "").strip()),
        destinatarios=destinatarios,
    )


def _corte(frecuencia: str) -> e.Corte:
    try:
        fecha, origen = resolver_corte(frecuencia, None)
        return e.Corte(fecha=fecha, origen=origen)
    except ConfiguracionError as exc:
        return e.Corte(fecha=None, origen=str(exc))


ENV = RAIZ / ".env"
# campo del pedido -> variable del .env. Los cortes se leen en cada ejecución; carpetas y driver se fijan al arrancar la API.
VARIABLES = {
    "corte_mensual": "FECHA_CORTE_MENSUAL", "corte_diario": "FECHA_CORTE_DIARIA",
    "dir_inputs": "REPORTES_DIR_INPUTS", "dir_outputs": "REPORTES_DIR_OUTPUTS", "driver_odbc": "DB_ODBC_DRIVER",
}
AL_REINICIAR = {
    "dir_inputs": ("carpeta de entradas", lambda v: Path(v) != DIR_INPUTS),
    "dir_outputs": ("carpeta de salidas", lambda v: Path(v) != DIR_OUTPUTS),
    "driver_odbc": ("driver ODBC", lambda v: v != DRIVER_ODBC),
}


def _pendientes_reinicio() -> list[str]:
    guardado = dotenv_values(ENV) if ENV.exists() else {}
    return [nombre for campo, (nombre, difiere) in AL_REINICIAR.items() if guardado.get(VARIABLES[campo]) and difiere(guardado[VARIABLES[campo]])]


def configuracion() -> e.ConfiguracionGeneral:
    return e.ConfiguracionGeneral(corte_mensual=_corte("mensual"), corte_diario=_corte("diaria"),
                                  dir_inputs=str(DIR_INPUTS), dir_outputs=str(DIR_OUTPUTS), driver_odbc=DRIVER_ODBC,
                                  pendientes_reinicio=_pendientes_reinicio())


def _validar_ajuste(campo: str, valor: str) -> str:
    if not valor:
        return ""
    if any(c in valor for c in "\r\n#\"'"):
        raise PeticionInvalida(f"{VARIABLES[campo]}: no admite saltos de línea, # ni comillas.")
    if campo.startswith("corte_"):
        try:
            corte = fecha_iso(valor)
            if campo == "corte_mensual":
                Cortes.mensual(corte)
        except (ValueError, ConfiguracionError) as exc:
            raise PeticionInvalida(f"{VARIABLES[campo]}: {exc}") from exc
        if corte > date.today():
            raise PeticionInvalida(f"{VARIABLES[campo]}: la fecha {valor} está en el futuro.")
    elif campo.startswith("dir_") and not Path(valor).is_absolute():
        raise PeticionInvalida(f"{VARIABLES[campo]}: escribe una ruta completa (con letra de unidad).")
    return valor


def actualizar_cuenta(actual: str, p: e.PedidoCuenta) -> str:
    """Cambia WEB_USUARIO y/o WEB_CLAVE en el .env (rige ya, sin reiniciar). Devuelve el usuario resultante."""
    try:
        if not sesion.iniciar(actual, p.clave_actual):  # misma verificación y límite de intentos que el inicio de sesión
            raise PeticionInvalida("La contraseña actual no es correcta.")
    except PermissionError as exc:
        raise PeticionInvalida(str(exc)) from exc
    usuario, clave = (p.usuario or "").strip(), p.clave_nueva or ""
    if not usuario and not clave:
        raise PeticionInvalida("No hay ningún cambio que guardar.")
    cambios: dict[str, str] = {}
    if usuario:
        if any(c in usuario for c in "\r\n#\"' \t"):
            raise PeticionInvalida("El usuario no admite espacios, saltos de línea, # ni comillas.")
        cambios["WEB_USUARIO"] = usuario
    if clave:
        if len(clave) < 8:
            raise PeticionInvalida("La contraseña nueva debe tener al menos 8 caracteres.")
        if any(c in clave for c in "\r\n#\"'"):
            raise PeticionInvalida("La contraseña no admite saltos de línea, # ni comillas (el .env no las soporta).")
        cambios["WEB_CLAVE"] = clave
    entorno.escribir(ENV, cambios)
    os.environ.update(cambios)
    return usuario or actual


def guardar_configuracion(p: e.PedidoConfiguracion) -> e.ConfiguracionGeneral:
    """Valida y escribe en el .env solo lo enviado. Los cortes rigen desde la próxima ejecución; carpetas y driver al reiniciar la API."""
    cambios = {VARIABLES[c]: _validar_ajuste(c, (getattr(p, c) or "").strip()) for c in p.model_fields_set}
    if not cambios:
        raise PeticionInvalida("No hay ningún ajuste que guardar.")
    entorno.escribir(ENV, cambios)
    for campo in ("corte_mensual", "corte_diario"):
        if campo in p.model_fields_set:  # las ejecuciones heredan os.environ: sin esto seguirían con el corte viejo
            if cambios[VARIABLES[campo]]:
                os.environ[VARIABLES[campo]] = cambios[VARIABLES[campo]]
            else:
                os.environ.pop(VARIABLES[campo], None)
    return configuracion()


def servidores() -> list[e.Servidor]:
    return [
        e.Servidor(nombre=s.nombre, servidor=s.servidor, autenticacion="SQL" if (s.usuario or not s.windows_auth_defecto) else "Windows",
                   credenciales_en_env=bool(s.usuario and s.clave), descripcion=s.descripcion, bases=list(BASES_DE[s.nombre]))
        for s in SERVIDORES.values()
    ]


def probar_servidor(nombre: str) -> e.PruebaConexion:
    """SELECT 1 contra un servidor (mismo camino que usan los reportes: db.leer_sql)."""
    if nombre not in SERVIDORES:
        raise NoEncontrado(f"Servidor desconocido: {nombre}")
    inicio = time.perf_counter()
    try:
        leer_sql(nombre, "SELECT 1 AS ok")
        ok, detalle = True, "Conexión correcta"
    except Exception as exc:  # noqa: BLE001 - diagnóstico para la pantalla, igual que probar_conexiones
        ok, detalle = False, f"{type(exc).__name__}: {str(exc)[:200]}"
    return e.PruebaConexion(nombre=nombre, ok=ok, detalle=detalle, milisegundos=round((time.perf_counter() - inicio) * 1000))


# ---------------------------------------------------------------- argumentos de ejecución

def argumentos(p: e.PedidoEjecucion) -> list[str]:
    """Valida el pedido con las reglas del reporte y devuelve los argumentos de `main.py`."""
    nombre = p.reporte
    corte = validar_corte(nombre, p.fecha_corte)
    lt = lote(nombre)
    if lt is None:
        if p.forzar or p.confirmar_escritura or p.correo == "todos":
            raise PeticionInvalida(f"«{nombre}» tiene lógica propia: no admite forzar, confirmar escritura ni envío de correo.")
        return [nombre, "--mes", f"{corte:%Y-%m}"] if nombre in USAN_MES else [nombre, "--fecha-corte", f"{corte:%Y-%m-%d}"]
    if lt.escribe_en_bd and not p.confirmar_escritura:
        raise PeticionInvalida(f"«{nombre}» crea/borra tablas permanentes: confirma la escritura en la base de datos.")
    args = [nombre, "--fecha-corte", f"{corte:%Y-%m-%d}"]
    if p.forzar:
        args.append("--forzar")
    if p.confirmar_escritura:
        args.append("--confirmar-escritura")
    if lt.entrega is not None:
        if p.correo == "todos":
            if not p.conforme:
                raise PeticionInvalida("Para enviar a toda la lista debes confirmar que revisaste el correo de prueba.")
            envio = estado_envio(nombre, corte)
            if not envio.prueba:
                raise PeticionInvalida(f"No hay correo de prueba para {corte:%Y-%m-%d}: primero ejecuta el reporte (envía la prueba).")
            if envio.todos:
                raise PeticionInvalida(f"Ya se envió a toda la lista el {envio.todos}.")
            args.append("--conforme")
        args += ["--correo", p.correo]
    elif p.correo == "todos":
        raise PeticionInvalida(f"«{nombre}» no envía correo.")
    return args


# ---------------------------------------------------------------- archivos generados y vista previa

def _carpeta(nombre: str) -> Path:
    reporte(nombre)
    return carpeta_salida(nombre)


def _ruta_segura(nombre: str, archivo: str) -> Path:
    """Solo archivos directamente dentro de la carpeta de salida del reporte (sin rutas relativas)."""
    carpeta = _carpeta(nombre).resolve()
    ruta = (carpeta / archivo).resolve()
    if ruta.parent != carpeta or not ruta.is_file():
        raise NoEncontrado(f"No existe el archivo {archivo}")
    return ruta


def archivos(nombre: str) -> list[e.Archivo]:
    carpeta = _carpeta(nombre)
    if not carpeta.is_dir():
        return []
    salida = [
        e.Archivo(nombre=f.name, tipo=TIPOS.get(f.suffix.lower(), "otro"), tamano=f.stat().st_size,
                  modificado=datetime.fromtimestamp(f.stat().st_mtime).isoformat(timespec="seconds"))
        for f in carpeta.iterdir() if f.is_file() and not f.name.startswith("~$")
    ]
    return sorted(salida, key=lambda a: a.modificado, reverse=True)


def ruta_descarga(nombre: str, archivo: str) -> Path:
    return _ruta_segura(nombre, archivo)


def _celda(v):
    if v is None or isinstance(v, (int, float, str)):
        return v
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v)


def vista_previa(nombre: str, archivo: str) -> e.VistaPrevia:
    ruta = _ruta_segura(nombre, archivo)
    tipo = TIPOS.get(ruta.suffix.lower())
    if tipo == "texto":
        return e.VistaPrevia(archivo=ruta.name, tipo="texto", texto=ruta.read_text(encoding="utf-8", errors="replace")[:MAX_TEXTO])
    if tipo != "excel":
        raise PeticionInvalida("Solo hay vista previa de Excel y texto; las imágenes se abren directamente.")
    from openpyxl import load_workbook

    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hojas = []
        for ws in libro.worksheets:
            filas = ws.iter_rows(values_only=True)
            columnas = [str(c) if c is not None else "" for c in next(filas, ())]
            datos = [[_celda(v) for v in fila] for _, fila in zip(range(FILAS_VISTA_PREVIA), filas)]
            hojas.append(e.HojaPrevia(nombre=ws.title, columnas=columnas, filas=datos, total_filas=max((ws.max_row or 1) - 1, 0)))
        return e.VistaPrevia(archivo=ruta.name, tipo="excel", hojas=hojas)
    finally:
        libro.close()


# ---------------------------------------------------------------- envío por correo (prueba -> conforme -> todos)

def estado_envio(nombre: str, fecha: date | None) -> e.EstadoEnvio:
    lt = lote(nombre)
    if lt is None or lt.entrega is None:
        raise NoEncontrado(f"«{nombre}» no envía correo.")
    corte = validar_corte(nombre, fecha)
    ruta = _carpeta(nombre) / f"estado_envio_{corte:%Y%m%d}.json"  # mismo archivo que escribe EntregaCorreoResumen
    datos = json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else {}
    excel = datos.get("excel")
    return e.EstadoEnvio(corte=corte, prueba=datos.get("prueba"), todos=datos.get("todos"), excel=Path(excel).name if excel else None)

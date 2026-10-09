"""Formato Excel del legado: tabla Medium2, resumen jerárquico (RESUMEN_v2), nombres de archivo y libro compartido."""

import warnings
from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from reportes.comun.excel import ESTILO_TABLA, FORMATO_MONEDA, Resumen, filas_jerarquicas, guardar_libro
from reportes.comun.fechas import nombre_de_archivo

NIVELES = ("NIVEL", "Grupo", "Territorio", "Corredor", "Agencia_Cli")
RESUMEN = Resumen("RESUMEN_v2", NIVELES, "Saldo_Capital", "TERRITORIO POR CLIENTE", "SALDO CARTERA - MIS")
LEGADO = Path(__file__).resolve().parents[1] / "data" / "inputs" / "cartera_sin_asignar_base.xlsm"


def _df():
    return pd.DataFrame({
        "NIVEL": ["FC"] * 5, "Grupo": ["GRUPAL", "GRUPAL", "GRUPAL", "INDIVIDUAL", "GRUPAL"],
        "Territorio": ["LIMA", "LIMA", "CENTRO", "NULL", "NULL"], "Corredor": ["LIMA ESTE", "LIMA ESTE", "PICHANAKI", None, None],
        "Agencia_Cli": ["AG ATE", "AG CHOSICA", "AG SATIPO", "OF ESPECIAL", "AG PANGOA"], "Saldo_Capital": [10.0, 5.5, 3.0, 2.0, 1.0],
    })


def test_nombres_de_archivo_del_legado():
    c = date(2026, 6, 30)
    assert nombre_de_archivo("Desembolsos_canal_{AAAAMMDD}", c) == "Desembolsos_canal_20260630"
    assert nombre_de_archivo("Datos Cierre {MES} {AA}", c) == "Datos Cierre Junio 26"
    assert nombre_de_archivo("Clientes_jóvenes_{mes3}{AA}", date(2026, 7, 31)) == "Clientes_jóvenes_jul26"
    assert nombre_de_archivo("clientes_nuevos_bancarizados_exclusivos_{MES3}{AA}", date(2026, 7, 31)) == "clientes_nuevos_bancarizados_exclusivos_JUL26"
    assert nombre_de_archivo("ClientesExtranjeros {MES3} {AAAA}-TIPODOC", date(2026, 7, 31)) == "ClientesExtranjeros JUL 2026-TIPODOC"


def test_hoja_de_datos_es_tabla_medium2_con_encabezados(tmp_path):
    ruta = guardar_libro(tmp_path / "x.xlsx", [("DATA_MIS_v2", _df())])
    ws = load_workbook(ruta)["DATA_MIS_v2"]
    assert [c.value for c in ws[1]] == list(_df().columns)
    (tabla,) = ws.tables.values()
    assert tabla.tableStyleInfo.name == ESTILO_TABLA == "TableStyleMedium2" and tabla.ref == "A1:F6"
    assert ws.max_row == 6


def test_hoja_sin_filas_conserva_encabezados_sin_tabla(tmp_path):
    ruta = guardar_libro(tmp_path / "x.xlsx", [("Castigos", pd.DataFrame(columns=["a", "b"]))])
    ws = load_workbook(ruta)["Castigos"]
    assert [c.value for c in ws[1]] == ["a", "b"] and not ws.tables


def test_resumen_jerarquico_sumas_orden_y_blanco():
    filas = filas_jerarquicas(_df(), NIVELES, "Saldo_Capital")
    assert filas[0] == (0, "FC", 21.5)
    assert (1, "GRUPAL", 19.5) in filas and (1, "INDIVIDUAL", 2.0) in filas
    grupal = [f for f in filas if f[0] == 2][:3]
    assert [f[1] for f in grupal] == ["CENTRO", "LIMA", "NULL"]          # A→Z, como la tabla dinámica
    assert (3, "(en blanco)", 1.0) in filas                               # corredor nulo


def test_hoja_resumen_con_formato_de_legado(tmp_path):
    ruta = guardar_libro(tmp_path / "x.xlsx", [("DATA_MIS_v2", _df())], [(RESUMEN, _df())])
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["DATA_MIS_v2", "RESUMEN_v2"]
    ws = wb["RESUMEN_v2"]
    assert (ws["B2"].value, ws["C2"].value) == ("TERRITORIO POR CLIENTE", "SALDO CARTERA - MIS")
    assert ws["B3"].value == "FC" and ws["C3"].number_format == FORMATO_MONEDA == '"S/" #,##0.00'
    assert [ws.cell(r, 2).alignment.indent for r in (3, 4, 5)] == [0, 1, 2]       # sangría por nivel
    ultimo = ws.max_row
    assert ws.cell(ultimo, 2).value == "Total general" and ws.cell(ultimo, 3).value == 21.5


def test_libro_compartido_conserva_las_otras_hojas(tmp_path):
    ruta = tmp_path / "Datos Cierre Junio 26.xlsx"
    guardar_libro(ruta, [("Captaciones", pd.DataFrame({"x": [1]}))], compartido=True)
    guardar_libro(ruta, [("Castigos", pd.DataFrame({"y": [2]}))], compartido=True)
    guardar_libro(ruta, [("Captaciones", pd.DataFrame({"x": [9]}))], compartido=True)   # se reemplaza solo esa hoja
    wb = load_workbook(ruta)
    assert set(wb.sheetnames) == {"Captaciones", "Castigos"}
    assert wb["Captaciones"]["A2"].value == 9 and wb["Castigos"]["A2"].value == 2


@pytest.mark.skipif(not LEGADO.exists(), reason="no está el Excel del legado")
def test_el_resumen_generado_coincide_con_el_RESUMEN_v2_del_legado():
    """Mismo orden, sangrías y sumas que la tabla dinámica del Excel original, calculado desde su DATA_MIS_v2."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(LEGADO, data_only=True)
    datos = list(wb["DATA_MIS_v2"].values)
    df = pd.DataFrame(datos[1:], columns=datos[0])
    esperado = [(wb["RESUMEN_v2"].cell(r, 2).alignment.indent, wb["RESUMEN_v2"].cell(r, 2).value, wb["RESUMEN_v2"].cell(r, 3).value)
                for r in range(3, wb["RESUMEN_v2"].max_row)]        # sin la fila «Total general»
    obtenido = filas_jerarquicas(df, NIVELES, "Saldo_Capital")
    redondeado = lambda filas: sorted((n, e, round(v, 2)) for n, e, v in filas)  # noqa: E731
    assert redondeado(obtenido) == redondeado([(int(n), e, v) for n, e, v in esperado])      # mismas filas y mismas sumas
    # Mismo orden en los niveles de agrupación; entre agencias hermanas la tabla dinámica del legado tiene un orden manual
    assert [(n, e) for n, e, _ in obtenido if n < 4] == [(int(n), e) for n, e, _ in esperado if n < 4]


def test_columnas_repetidas_o_sin_nombre_no_rompen_la_exportacion(tmp_path):
    df = pd.DataFrame([[1, 2, 3]], columns=["a", "a", None])
    ws = load_workbook(guardar_libro(tmp_path / "x.xlsx", [("H", df)]))["H"]
    assert [c.value for c in ws[1]] == ["a", "a_2", "Columna"]
